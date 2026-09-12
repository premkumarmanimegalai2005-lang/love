from fastapi import FastAPI
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from openpyxl import Workbook, load_workbook
from datetime import datetime
import os

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

EXCEL_FILE = "demo_results.xlsx"


class UserData(BaseModel):
    username: str
    password: str


def create_excel():
    if not os.path.exists(EXCEL_FILE):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Users"

        sheet.append([
            "Username",
            "Password",
            "Created Time"
        ])

        workbook.save(EXCEL_FILE)


@app.post("/register")
def register(data: UserData):

    create_excel()

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Users"]

    # Check whether username already exists
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == data.username:
            workbook.close()

            return {
                "status": "Failed",
                "message": "Username already exists"
            }

    # Create new account
    sheet.append([
        data.username,
        data.password,
        datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ])

    workbook.save(EXCEL_FILE)
    workbook.close()

    return {
        "status": "Success",
        "message": "Account created successfully"
    }


@app.post("/login")
def login(data: UserData):

    create_excel()

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Users"]

    for row in sheet.iter_rows(min_row=2, values_only=True):

        username = row[0]
        password = row[1]

        if username == data.username and password == data.password:

            workbook.close()

            return {
                "status": "Success",
                "message": "Login successful"
            }

    workbook.close()

    return {
        "status": "Failed",
        "message": "Invalid username or password"
    }